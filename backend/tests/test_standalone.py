import base64
import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from banking_crosswalk.app import app
from banking_crosswalk.catalog import load_aiaf_question_bank
from banking_crosswalk.policy_documents import extract_policy_file


def test_independent_health_and_synthetic_metadata():
    with TestClient(app) as client:
        assert client.get("/health").json()["version"] == "0.1.0"
        metadata = client.get("/v1/demo").json()
        assert metadata["questionnaire"]["source_version"] == "SYNTHETIC-DEMO-1.0"
        assert metadata["questionnaire"]["focus_question_count"] == 3
        assert metadata["questionnaire"]["core_question_count"] == 16


def test_inline_coverage_is_traceable_and_never_auto_approved():
    response = TestClient(app).post(
        "/v1/assessments/coverage",
        headers={"X-Tenant-ID": "standalone"},
        json={
            "use_case": {"apra_regulated": True},
            "questionnaire_pack": "DEMO_BANKING_V1",
            "persist": False,
        },
    )
    assert response.status_code == 201
    result = response.json()
    assert result["questionnaire_sha256"]
    assert result["metrics"]["confirmed_coverage_rate"] == 0
    assert result["metrics"]["applicability_undetermined"] > 0
    assert result["gaps"]


def test_request_body_limit_checks_declared_length():
    response = TestClient(app).post(
        "/v1/assessments/coverage",
        headers={"Content-Length": str(33 * 1024 * 1024)},
        content=b"{}",
    )
    assert response.status_code == 413


def test_streamed_request_without_length_cannot_bypass_byte_limit(monkeypatch):
    import banking_crosswalk.app as application

    monkeypatch.setattr(application, "MAX_REQUEST_BYTES", 5)
    response = TestClient(app).post(
        "/v1/assessments/coverage", content=iter([b"123", b"456"])
    )
    assert response.status_code == 413


def test_policy_docx_extraction_and_api_import():
    document = Document()
    document.add_paragraph(
        "Refunds above AUD 500 require human approval before payment execution."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    content = buffer.getvalue()
    assert "AUD 500" in extract_policy_file("policy.docx", content)
    response = TestClient(app).post(
        "/v1/policies/import-file",
        headers={"X-Tenant-ID": "file-import"},
        json={
            "name": "Synthetic authority policy",
            "filename": "policy.docx",
            "file_base64": base64.b64encode(content).decode(),
        },
    )
    assert response.status_code == 201
    assert response.json()["clauses"][0]["thresholds_aud"] == [500]


def test_policy_pdf_and_plain_text_parsing():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    assert isinstance(extract_policy_file("empty.pdf", buffer.getvalue()), str)
    assert (
        extract_policy_file("policy.txt", b"Human approval is required.")
        == "Human approval is required."
    )


def test_docx_rejects_entity_expansion():
    document = Document()
    raw = io.BytesIO()
    document.save(raw)
    modified = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source,
        zipfile.ZipFile(modified, "w") as target,
    ):
        for member in source.infolist():
            data = source.read(member)
            if member.filename == "word/document.xml":
                data = (
                    b'<!DOCTYPE doc [<!ENTITY attack "expanded">]><doc>&attack;</doc>'
                )
            target.writestr(member.filename, data)
    with pytest.raises(Exception):
        extract_policy_file("malicious.docx", modified.getvalue())


def test_malformed_policy_file_api_does_not_expose_parser_errors():
    response = TestClient(app).post(
        "/v1/policies/import-file",
        headers={"X-Tenant-ID": "bad-file"},
        json={
            "name": "Bad document",
            "filename": "bad.pdf",
            "file_base64": base64.b64encode(b"not a pdf").decode(),
        },
    )
    assert response.status_code == 422
    assert (
        response.json()["detail"]
        == "Policy file could not be extracted within supported limits"
    )


def test_core_module_has_no_parent_project_catalog_dependency():
    bank = load_aiaf_question_bank()
    assert bank.metadata.source_filename == "demo_questionnaire.json"
    assert len(bank.questions) == 19


def test_exact_multiple_policy_thresholds_cannot_be_partially_matched():
    from banking_crosswalk.policy import _questionnaire_covers_generic_facet
    from banking_crosswalk.models import PolicyClause

    bank = load_aiaf_question_bank()
    bank.questions[0].question_text = "Does a refund above AUD 500 require approval?"
    clause = PolicyClause(
        clause_id="multi",
        source_name="Synthetic",
        text="Approval bands AUD 500 and AUD 1000.",
        mandatory=True,
        thresholds_aud=[500, 1000],
    )
    assert not _questionnaire_covers_generic_facet(
        "approval_threshold_defined", bank, clause
    )
